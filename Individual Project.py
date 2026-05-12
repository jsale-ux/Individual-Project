import tkinter as tk
from tkinter import messagebox
import os
from datetime import datetime
from openpyxl import Workbook, load_workbook

# This finds the folder where the Python file is saved.
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# This is the Excel file where records will be stored.
EXCEL_FILE = os.path.join(BASE_DIR, "autocare_records.xlsx")

# These are the column names used in the Excel sheet.
HEADERS = [
    "date",
    "miles",
    "gas_cost",
    "repair_cost",
    "service_type",
    "cost_per_mile",
    "maintenance_status"
]


def create_excel_file():
    # This function creates the Excel file if it does not exist.
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Records"
        ws.append(HEADERS)
        wb.save(EXCEL_FILE)
    else:
        # This checks whether the first row has the correct headers.
        try:
            wb = load_workbook(EXCEL_FILE)
            ws = wb.active

            first_row = [cell.value for cell in ws[1]]

            if first_row != HEADERS:
                ws.delete_rows(1, ws.max_row)
                ws.append(HEADERS)
                wb.save(EXCEL_FILE)
        except Exception:
            wb = Workbook()
            ws = wb.active
            ws.title = "Records"
            ws.append(HEADERS)
            wb.save(EXCEL_FILE)


def save_to_excel(data):
    # This function adds one new row of data into the Excel file.
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    ws.append(data)
    wb.save(EXCEL_FILE)


def calculate_and_save():
    # This function gets the user's input from the boxes.
    date_text = date_entry.get().strip()
    miles_text = miles_entry.get().strip()
    gas_text = gas_entry.get().strip()
    repair_text = repair_entry.get().strip()
    service_text = service_entry.get().strip()

    # This checks if any box is empty.
    if not date_text or not miles_text or not gas_text or not repair_text or not service_text:
        messagebox.showerror("Missing Information", "Please fill in all boxes.")
        return

    try:
        # These lines change text into numbers so math can be done.
        miles = float(miles_text)
        gas_cost = float(gas_text)
        repair_cost = float(repair_text)

        # This checks if the date is written correctly.
        datetime.strptime(date_text, "%Y-%m-%d")
    except ValueError:
        messagebox.showerror("Invalid Input", "Use numbers and date format YYYY-MM-DD.")
        return

    # Miles cannot be zero or negative.
    if miles <= 0:
        messagebox.showerror("Invalid Miles", "Miles must be greater than 0.")
        return

    # This calculates the cost per mile.
    cost_per_mile = (gas_cost + repair_cost) / miles

    # This creates a simple maintenance alert level.
    if repair_cost >= 300 or miles >= 5000:
        status = "High"
    elif repair_cost >= 100 or miles >= 3000:
        status = "Medium"
    else:
        status = "Low"

    # This saves the information into the Excel file.
    save_to_excel([
        date_text,
        round(miles, 2),
        round(gas_cost, 2),
        round(repair_cost, 2),
        service_text,
        round(cost_per_mile, 4),
        status
    ])

    # This shows the result to the user.
    result_label.config(
        text=f"Saved! Cost per mile: ${cost_per_mile:.2f} | Alert: {status}"
    )

    # This updates the summary section.
    show_summary()

    # These lines clear some boxes for the next entry.
    miles_entry.delete(0, tk.END)
    gas_entry.delete(0, tk.END)
    repair_entry.delete(0, tk.END)
    service_entry.delete(0, tk.END)
    service_entry.insert(0, "Oil Change")


def show_summary():
    # This function reads the Excel file and shows totals.
    create_excel_file()

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    # This gets all rows except the header row.
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    # If there are no saved records yet, show this message.
    if not rows:
        summary_label.config(text="No records saved yet.")
        return

    # These lines add up the numbers from all saved rows.
    total_miles = sum(float(row[1] or 0) for row in rows)
    total_gas = sum(float(row[2] or 0) for row in rows)
    total_repairs = sum(float(row[3] or 0) for row in rows)

    # This shows the totals in the summary label.
    summary_label.config(
        text=(
            f"Total Records: {len(rows)}\n"
            f"Total Miles: {total_miles:.2f}\n"
            f"Total Gas Cost: ${total_gas:.2f}\n"
            f"Total Repair Cost: ${total_repairs:.2f}"
        )
    )


# This creates the main window.
window = tk.Tk()
window.title("Simple AutoCare Tracker")
window.geometry("430x460")
window.configure(bg="#f2f2f2")

# This makes sure the Excel file exists when the app starts.
create_excel_file()

# This is the title at the top of the window.
title_label = tk.Label(
    window,
    text="Simple AutoCare Tracker",
    font=("Arial", 16, "bold"),
    bg="#f2f2f2"
)
title_label.pack(pady=12)

# This frame holds all the labels and input boxes.
form_frame = tk.Frame(window, bg="#ffffff", padx=15, pady=15)
form_frame.pack(padx=15, pady=10, fill="both")

# These labels and boxes collect user input.
tk.Label(form_frame, text="Date (YYYY-MM-DD)", bg="#ffffff").grid(row=0, column=0, sticky="w", pady=6, padx=5)
date_entry = tk.Entry(form_frame, width=25)
date_entry.grid(row=0, column=1, pady=6, padx=5)
date_entry.insert(0, datetime.today().strftime("%Y-%m-%d"))

tk.Label(form_frame, text="Miles Driven", bg="#ffffff").grid(row=1, column=0, sticky="w", pady=6, padx=5)
miles_entry = tk.Entry(form_frame, width=25)
miles_entry.grid(row=1, column=1, pady=6, padx=5)

tk.Label(form_frame, text="Gas Cost ($)", bg="#ffffff").grid(row=2, column=0, sticky="w", pady=6, padx=5)
gas_entry = tk.Entry(form_frame, width=25)
gas_entry.grid(row=2, column=1, pady=6, padx=5)

tk.Label(form_frame, text="Repair Cost ($)", bg="#ffffff").grid(row=3, column=0, sticky="w", pady=6, padx=5)
repair_entry = tk.Entry(form_frame, width=25)
repair_entry.grid(row=3, column=1, pady=6, padx=5)

tk.Label(form_frame, text="Service Type", bg="#ffffff").grid(row=4, column=0, sticky="w", pady=6, padx=5)
service_entry = tk.Entry(form_frame, width=25)
service_entry.grid(row=4, column=1, pady=6, padx=5)
service_entry.insert(0, "Oil Change")

# This button runs the calculation and saves the data.
save_button = tk.Button(
    window,
    text="Calculate and Save",
    font=("Arial", 11, "bold"),
    command=calculate_and_save,
    bg="#4c7ef3",
    fg="white",
    width=18
)
save_button.pack(pady=10)

# This label shows the result after saving.
result_label = tk.Label(
    window,
    text="",
    bg="#f2f2f2",
    font=("Arial", 10)
)
result_label.pack(pady=8)

# This title shows the summary section.
summary_title = tk.Label(
    window,
    text="Summary",
    font=("Arial", 12, "bold"),
    bg="#f2f2f2"
)
summary_title.pack(pady=(10, 0))

# This label displays the total saved information.
summary_label = tk.Label(
    window,
    text="",
    bg="#f2f2f2",
    font=("Arial", 10),
    justify="left"
)
summary_label.pack(pady=8)

# This loads the summary when the app first opens.
show_summary()

# This keeps the program running.
window.mainloop()